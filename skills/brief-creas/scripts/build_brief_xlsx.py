#!/usr/bin/env python3
"""
Génère le brief créas en xlsx (3 onglets) à partir d'un JSON.

Usage:
    python build_brief_xlsx.py brief.json sortie.xlsx

Format du JSON attendu:

{
  "client": "Nom du client",
  "concepts": [
    {
      "n": 1,
      "axe": "Revenu",
      "concept": "Le chiffre plein écran",
      "message": "Texte exact à afficher sur le visuel",
      "visuel": "Principe visuel décrit au DA",
      "traitement": "Ton et traitement graphique",
      "format": "4:5",
      "assets": "Assets nécessaires",
      "priorite": "P1",
      "inspi": "https://...",
      "note": "Note libre, mention des paires de test A/B"
    }
  ],
  "regles": [
    "Un seul message par visuel...",
    "..."
  ],
  "inspi": [
    {"source": "Nom", "lien": "https://...", "usage": "Ce qu'on y cherche et sa limite"}
  ],
  "assets": [
    {"asset": "Charte graphique", "concepts": "Tous", "statut": "A recevoir", "bloquant": "Oui"}
  ]
}

Seul "concepts" est obligatoire. Les autres sections sont générées si présentes.
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"

HDR_FILL = PatternFill("solid", start_color="16213E")
P_FILLS = {
    "P1": PatternFill("solid", start_color="FDE7E9"),
    "P2": PatternFill("solid", start_color="FFF4E0"),
    "P3": PatternFill("solid", start_color="EAF3EA"),
}
THIN = Side(style="thin", color="D6D6D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
CELL_FONT = Font(name=FONT, size=10)
BOLD_FONT = Font(name=FONT, size=10, bold=True)
LINK_FONT = Font(name=FONT, size=10, color="0563C1", underline="single")
TITLE_FONT = Font(name=FONT, size=13, bold=True, color="16213E")


def style_header(ws, row=1):
    for cell in ws[row]:
        if cell.value is not None:
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
            cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
            cell.border = BORDER
    ws.row_dimensions[row].height = 30


def set_widths(ws, widths):
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def sheet_concepts(wb, concepts):
    ws = wb.active
    ws.title = "Brief creas"
    headers = [
        "N°", "Axe", "Concept", "Message a afficher (texte exact)", "Principe visuel",
        "Traitement / ton", "Format prio", "Assets requis", "Priorite", "Inspi (lien)", "Note",
    ]
    ws.append(headers)

    keys = ["n", "axe", "concept", "message", "visuel",
            "traitement", "format", "assets", "priorite", "inspi", "note"]
    for c in concepts:
        ws.append([c.get(k, "") for k in keys])

    style_header(ws)

    for i in range(2, len(concepts) + 2):
        priorite = ws.cell(row=i, column=9).value
        for j in range(1, len(headers) + 1):
            cell = ws.cell(row=i, column=j)
            cell.font = CELL_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if j == 2:
                cell.font = BOLD_FONT
            if j == 9:
                cell.fill = P_FILLS.get(priorite, P_FILLS["P3"])
                cell.font = BOLD_FONT
                cell.alignment = Alignment(vertical="top", horizontal="center")
        link = ws.cell(row=i, column=10)
        if link.value and str(link.value).startswith("http"):
            link.hyperlink = link.value
            link.font = LINK_FONT
        ws.row_dimensions[i].height = 92

    set_widths(ws, [5, 16, 24, 52, 50, 30, 10, 32, 10, 42, 46])
    ws.freeze_panes = "D2"


def sheet_regles(wb, regles, inspi):
    ws = wb.create_sheet("Regles & inspi")
    ws["A1"] = "Regles de production (a lire avant de commencer)"
    ws["A1"].font = TITLE_FONT
    ws.append([])
    ws.append(["#", "Regle"])
    for idx, regle in enumerate(regles, start=1):
        ws.append([str(idx), regle])
    style_header(ws, row=3)

    for i in range(4, 4 + len(regles)):
        for j in (1, 2):
            cell = ws.cell(row=i, column=j)
            cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[i].height = 34

    start = 4 + len(regles) + 2
    ws.cell(row=start, column=1, value="Sources d'inspiration").font = TITLE_FONT
    ws.cell(row=start + 1, column=1, value="Source")
    ws.cell(row=start + 1, column=2, value="Lien")
    ws.cell(row=start + 1, column=3, value="Ce qu'on y cherche / limite")
    style_header(ws, row=start + 1)

    row = start + 2
    for item in inspi:
        ws.cell(row=row, column=1, value=item.get("source", "")).font = BOLD_FONT
        link = ws.cell(row=row, column=2, value=item.get("lien", ""))
        if str(link.value).startswith("http"):
            link.hyperlink = link.value
            link.font = LINK_FONT
        else:
            link.font = CELL_FONT
        ws.cell(row=row, column=3, value=item.get("usage", "")).font = CELL_FONT
        for j in (1, 2, 3):
            cell = ws.cell(row=row, column=j)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 46
        row += 1

    set_widths(ws, [26, 46, 88])


def sheet_assets(wb, assets):
    ws = wb.create_sheet("Assets a fournir")
    ws["A1"] = "Assets a demander au client"
    ws["A1"].font = TITLE_FONT
    ws.append([])
    ws.append(["Asset", "Pour quels concepts", "Statut", "Bloquant ?"])
    style_header(ws, row=3)

    for item in assets:
        ws.append([
            item.get("asset", ""), item.get("concepts", ""),
            item.get("statut", ""), item.get("bloquant", ""),
        ])

    for i in range(4, 4 + len(assets)):
        for j in range(1, 5):
            cell = ws.cell(row=i, column=j)
            cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if j == 4 and str(cell.value).strip().lower() == "oui":
                cell.fill = P_FILLS["P1"]
                cell.font = BOLD_FONT
        ws.row_dimensions[i].height = 32

    set_widths(ws, [62, 34, 30, 14])


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)

    with open(sys.argv[1], encoding="utf-8") as fh:
        data = json.load(fh)

    concepts = data.get("concepts", [])
    if not concepts:
        print("Erreur : le JSON doit contenir au moins un concept.")
        sys.exit(1)

    wb = Workbook()
    sheet_concepts(wb, concepts)
    if data.get("regles") or data.get("inspi"):
        sheet_regles(wb, data.get("regles", []), data.get("inspi", []))
    if data.get("assets"):
        sheet_assets(wb, data["assets"])

    wb.save(sys.argv[2])
    print(f"OK — {len(concepts)} concepts écrits dans {sys.argv[2]}")


if __name__ == "__main__":
    main()
