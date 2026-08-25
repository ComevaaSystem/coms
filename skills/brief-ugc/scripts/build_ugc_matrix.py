#!/usr/bin/env python3
"""
Génère la feuille de tournage UGC et la matrice de combinaisons en xlsx.

Usage:
    python build_ugc_matrix.py scripts.json sortie.xlsx [--max-combi 20]

Format du JSON attendu:

{
  "client": "Nom du client",
  "intervenant": "Thomas, DG",
  "themes": [
    {
      "nom": "Revenu",
      "mots_cles": ["50 a 100 euros de marge", "20 minutes"],
      "hooks": [
        "Texte du hook 1",
        "Texte du hook 2"
      ],
      "corps": [
        "Le principe. Description en une ou deux phrases.",
        "Le benefice concret. Description."
      ],
      "cta": [
        "Texte du CTA 1"
      ],
      "inserts": [
        "Mains qui clipsent le boitier, plan serre"
      ]
    }
  ]
}

Seuls "themes" avec "nom", "hooks", "corps" et "cta" sont obligatoires.

Identifiants générés: REVENU-H1, REVENU-C1, REVENU-A1.
Références de variantes: REVENU-H1-C1-A1, directement utilisables comme nom
de créa dans Meta ou TikTok pour lire la performance élément par élément.
"""

import argparse
import json
import re
import sys
from itertools import product

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"

HDR_FILL = PatternFill("solid", start_color="16213E")
TYPE_FILLS = {
    "Hook": PatternFill("solid", start_color="FDE7E9"),
    "Corps": PatternFill("solid", start_color="FFF4E0"),
    "CTA": PatternFill("solid", start_color="EAF3EA"),
}
THIN = Side(style="thin", color="D6D6D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
CELL_FONT = Font(name=FONT, size=10)
BOLD_FONT = Font(name=FONT, size=10, bold=True)
MONO_FONT = Font(name="Courier New", size=10, bold=True)
TITLE_FONT = Font(name=FONT, size=13, bold=True, color="16213E")

DUREES = {"Hook": "3 a 5 sec", "Corps": "15 a 20 sec", "CTA": "3 a 5 sec"}


def slug(nom):
    """Revenu -> REVENU, Preuve sociale -> PREUVE."""
    base = re.sub(r"[^A-Za-z0-9 ]", "", nom).strip().upper()
    return base.split(" ")[0][:10] or "THEME"


def style_header(ws, row=1):
    for cell in ws[row]:
        if cell.value is not None:
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
            cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
            cell.border = BORDER
    ws.row_dimensions[row].height = 28


def set_widths(ws, widths):
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def build_elements(themes):
    """Retourne {theme_nom: {"H": [(id, texte)], "C": [...], "A": [...]}}."""
    elements = {}
    for theme in themes:
        code = slug(theme["nom"])
        elements[theme["nom"]] = {
            "code": code,
            "H": [(f"{code}-H{i}", t) for i, t in enumerate(theme.get("hooks", []), 1)],
            "C": [(f"{code}-C{i}", t) for i, t in enumerate(theme.get("corps", []), 1)],
            "A": [(f"{code}-A{i}", t) for i, t in enumerate(theme.get("cta", []), 1)],
        }
    return elements


def sheet_tournage(wb, themes, elements, intervenant):
    ws = wb.active
    ws.title = "Tournage"
    ws["A1"] = f"Feuille de tournage — {intervenant}" if intervenant else "Feuille de tournage"
    ws["A1"].font = TITLE_FONT
    ws.append([])
    ws.append(["ID", "Theme", "Type", "Texte / points a faire passer", "Duree cible", "Prises", "Fait"])
    style_header(ws, row=3)

    row = 4
    for theme in themes:
        nom = theme["nom"]
        pack = elements[nom]
        for kind, label in (("H", "Hook"), ("C", "Corps"), ("A", "CTA")):
            for eid, texte in pack[kind]:
                ws.cell(row=row, column=1, value=eid).font = MONO_FONT
                ws.cell(row=row, column=2, value=nom).font = BOLD_FONT
                ws.cell(row=row, column=3, value=label)
                ws.cell(row=row, column=4, value=texte)
                ws.cell(row=row, column=5, value=DUREES[label])
                ws.cell(row=row, column=6, value="")
                ws.cell(row=row, column=7, value="")
                for j in range(1, 8):
                    cell = ws.cell(row=row, column=j)
                    cell.border = BORDER
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if j not in (1, 2):
                        cell.font = CELL_FONT
                    if j == 3:
                        cell.fill = TYPE_FILLS[label]
                        cell.font = BOLD_FONT
                ws.row_dimensions[row].height = 46
                row += 1

    set_widths(ws, [16, 20, 12, 86, 14, 10, 8])
    ws.freeze_panes = "A4"

    legend = row + 1
    ws.cell(row=legend, column=1,
            value=("Chaque element est tourne separement, en 2 ou 3 prises. "
                   "Aucun element ne doit faire reference a un autre. "
                   "Cocher la colonne Fait au fur et a mesure."))
    ws.cell(row=legend, column=1).font = Font(name=FONT, size=9, italic=True)
    ws.merge_cells(start_row=legend, start_column=1, end_row=legend, end_column=7)
    ws.cell(row=legend, column=1).alignment = Alignment(vertical="top", wrap_text=True)


def sheet_combinaisons(wb, themes, elements, max_combi):
    ws = wb.create_sheet("Combinaisons")
    ws.append(["Ref variante", "Theme", "ID hook", "ID corps", "ID CTA", "Texte du hook", "Statut"])
    style_header(ws)

    row = 2
    total = 0
    for theme in themes:
        nom = theme["nom"]
        pack = elements[nom]
        if not (pack["H"] and pack["C"] and pack["A"]):
            continue
        combis = list(product(pack["H"], pack["C"], pack["A"]))
        total += len(combis)
        for hook, corps, cta in combis[:max_combi]:
            ref = f"{pack['code']}-{hook[0].split('-')[1]}-{corps[0].split('-')[1]}-{cta[0].split('-')[1]}"
            ws.cell(row=row, column=1, value=ref).font = MONO_FONT
            ws.cell(row=row, column=2, value=nom).font = BOLD_FONT
            ws.cell(row=row, column=3, value=hook[0]).font = MONO_FONT
            ws.cell(row=row, column=4, value=corps[0]).font = MONO_FONT
            ws.cell(row=row, column=5, value=cta[0]).font = MONO_FONT
            ws.cell(row=row, column=6, value=hook[1]).font = CELL_FONT
            ws.cell(row=row, column=7, value="A monter").font = CELL_FONT
            for j in range(1, 8):
                cell = ws.cell(row=row, column=j)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[row].height = 34
            row += 1

    set_widths(ws, [24, 18, 16, 16, 16, 72, 14])
    ws.freeze_panes = "A2"

    note = row + 1
    ws.cell(row=note, column=1,
            value=(f"{total} combinaisons possibles au total, {row - 2} listees ici "
                   "(plafond par theme). La ref de variante sert de nom de crea dans Meta "
                   "ou TikTok : c'est ce qui permet de lire la performance hook par hook."))
    ws.cell(row=note, column=1).font = Font(name=FONT, size=9, italic=True)
    ws.merge_cells(start_row=note, start_column=1, end_row=note, end_column=7)
    ws.cell(row=note, column=1).alignment = Alignment(vertical="top", wrap_text=True)


def sheet_inserts(wb, themes):
    rows = [(t["nom"], i) for t in themes for i in t.get("inserts", [])]
    if not rows:
        return
    ws = wb.create_sheet("Inserts")
    ws.append(["Theme", "Insert / plan d'illustration", "Statut"])
    style_header(ws)
    for idx, (theme, insert) in enumerate(rows, start=2):
        ws.cell(row=idx, column=1, value=theme).font = BOLD_FONT
        ws.cell(row=idx, column=2, value=insert).font = CELL_FONT
        ws.cell(row=idx, column=3, value="A produire").font = CELL_FONT
        for j in range(1, 4):
            cell = ws.cell(row=idx, column=j)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[idx].height = 32
    set_widths(ws, [20, 78, 16])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("sortie")
    parser.add_argument("--max-combi", type=int, default=20,
                        help="Nombre max de combinaisons listees par theme")
    args = parser.parse_args()

    with open(args.source, encoding="utf-8") as fh:
        data = json.load(fh)

    themes = data.get("themes", [])
    if not themes:
        print("Erreur : le JSON doit contenir au moins un theme.")
        sys.exit(1)

    elements = build_elements(themes)

    wb = Workbook()
    sheet_tournage(wb, themes, elements, data.get("intervenant", ""))
    sheet_combinaisons(wb, themes, elements, args.max_combi)
    sheet_inserts(wb, themes)
    wb.save(args.sortie)

    n_h = sum(len(e["H"]) for e in elements.values())
    n_c = sum(len(e["C"]) for e in elements.values())
    n_a = sum(len(e["A"]) for e in elements.values())
    print(f"OK — {len(themes)} themes, {n_h} hooks, {n_c} corps, {n_a} CTA -> {args.sortie}")


if __name__ == "__main__":
    main()
